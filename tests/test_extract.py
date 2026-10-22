import pytest

from specdiff_ai.extract import load_document, parse_lines

SAMPLE = [
    "电阻点焊工艺技术协议",
    "1 总则",
    "1.1 本协议适用于车身电阻点焊工艺。",
    "1.2 设备应满足节拍要求。",
    "2 焊接参数",
    "2.1 焊接电流应为 8.5 kA。",
    "2.2 电极压力为 3.2 kN。",
]


def test_parse_lines_sections_and_clauses():
    clauses = parse_lines(SAMPLE, "sample.txt")
    assert len(clauses) == 5  # title line becomes an unnumbered clause
    numbered = {c.clause_no: c for c in clauses if c.clause_no}
    assert set(numbered) == {"1.1", "1.2", "2.1", "2.2"}
    assert numbered["1.1"].section_path == "1 总则"
    assert numbered["2.1"].section_path == "2 焊接参数"
    assert numbered["2.1"].params[0].value == 8.5
    assert numbered["2.1"].locator == "第6行"


def test_markdown_headings():
    clauses = parse_lines(
        ["# 协议", "## 参数要求", "电流为 8.5 kA。"], "a.md"
    )
    assert len(clauses) == 1
    assert clauses[0].section_path == "协议 > 参数要求"


def test_chinese_chapter_headings():
    clauses = parse_lines(
        ["第一章 总则", "本协议适用范围为全部焊点。", "第二章 参数", "电流 8.5 kA。"],
        "a.txt",
    )
    assert clauses[0].section_path == "第一章 总则"
    assert clauses[1].section_path == "第二章 参数"


def test_load_txt_file(tmp_path):
    path = tmp_path / "old.txt"
    path.write_text("\n".join(SAMPLE), encoding="utf-8")
    clauses = load_document(path)
    assert len(clauses) == 5
    assert clauses[0].doc == "old.txt"


def test_load_docx_file(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("焊接参数", level=1)
    document.add_paragraph("2.1 焊接电流应为 8.5 kA。")
    document.add_paragraph("2.2 电极压力为 3.2 kN。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "检验项"
    table.rows[0].cells[1].text = "焊核直径 4.5 mm"
    path = tmp_path / "spec.docx"
    document.save(str(path))

    clauses = load_document(path)
    by_no = {c.clause_no: c for c in clauses if c.clause_no}
    assert by_no["2.1"].section_path == "焊接参数"
    assert by_no["2.1"].params[0].unit == "kA"
    assert any("焊核直径" in c.text for c in clauses)  # table row flattened


def _make_pdf(lines):
    def esc(text):
        return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    content = "BT /F1 12 Tf 72 720 Td " + " ".join(
        "(%s) Tj 0 -16 Td " % esc(line) for line in lines
    ) + "ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = "%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf.encode("latin-1")))
        pdf += "%d 0 obj\n%s\nendobj\n" % (i, obj)
    xref_pos = len(pdf.encode("latin-1"))
    pdf += "xref\n0 %d\n" % (len(objects) + 1)
    pdf += "0000000000 65535 f \n"
    for offset in offsets:
        pdf += "%010d 00000 n \n" % offset
    pdf += "trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_pos,
    )
    return pdf.encode("latin-1")


def test_load_pdf_file(tmp_path):
    path = tmp_path / "spec.pdf"
    path.write_bytes(
        _make_pdf(["1 General", "1.1 Welding current shall be 8.5 kA."])
    )
    clauses = load_document(path)
    assert len(clauses) == 1
    assert clauses[0].clause_no == "1.1"
    assert clauses[0].params[0].value == 8.5
    assert clauses[0].locator.startswith("第1页")


def test_unsupported_suffix(tmp_path):
    path = tmp_path / "spec.xlsx"
    path.write_bytes(b"nope")
    with pytest.raises(ValueError):
        load_document(path)
