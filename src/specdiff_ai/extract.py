"""Document loading and clause extraction.

Supported inputs:

* ``.docx`` via python-docx (heading styles drive the section tree;
  table rows are flattened into clauses)
* ``.pdf`` via pypdf (plain text layer only -- scanned PDFs need OCR,
  which is out of scope for V1)
* ``.txt`` / ``.md`` / ``.text`` via the standard library

Every extracted :class:`Clause` records where it came from (document
label + locator such as ``第3页`` / ``第12行`` / ``第5段``) so that any
diff conclusion can be traced back to the original wording.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Sequence, Tuple

from .params import Param, extract_params

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt", ".md", ".text"}

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_CN_HEADING_RE = re.compile(r"^第[一二三四五六七八九十百零\d]+(章|节)\s*(.*)$")
_NUMBERED_RE = re.compile(r"^(\d+(?:\.\d+)*)[\.、]?\s+(.+)$")


@dataclass
class Clause:
    """One numbered (or unnumbered) clause of a technical agreement."""

    doc: str
    section_path: str
    clause_no: str
    text: str
    locator: str
    params: List[Param] = field(default_factory=list)

    def __post_init__(self) -> None:
        if not self.params:
            self.params = extract_params(self.text)

    @property
    def key(self) -> str:
        return "%s|%s" % (self.section_path, self.clause_no)


def _looks_like_heading(number: str, rest: str) -> bool:
    """A short numbered line without sentence punctuation is a heading."""
    depth = number.count(".") + 1
    if depth > 2:
        return False
    if len(rest) > 40:
        return False
    if rest.endswith(("。", "．", ".", "；", ";", "，", ",")):
        return False
    return True


def parse_lines(
    lines: Sequence[str], doc_label: str, locator_fmt: str = "第%d行"
) -> List[Clause]:
    """Turn raw text lines into clauses with a section tree."""
    clauses: List[Clause] = []
    section_stack: List[Tuple[int, str]] = []  # (depth, title)

    def current_path() -> str:
        return " > ".join(title for _, title in section_stack)

    def push_heading(depth: int, title: str) -> None:
        while section_stack and section_stack[-1][0] >= depth:
            section_stack.pop()
        section_stack.append((depth, title))

    for index, raw in enumerate(lines, start=1):
        line = raw.strip()
        if not line:
            continue
        locator = locator_fmt % index

        md = _MD_HEADING_RE.match(line)
        if md:
            push_heading(len(md.group(1)), md.group(2).strip())
            continue

        cn = _CN_HEADING_RE.match(line)
        if cn:
            push_heading(1 if cn.group(1) == "章" else 2, line)
            continue

        numbered = _NUMBERED_RE.match(line)
        if numbered:
            number, rest = numbered.group(1), numbered.group(2).strip()
            if _looks_like_heading(number, rest):
                push_heading(number.count(".") + 1, line)
                continue
            clauses.append(
                Clause(
                    doc=doc_label,
                    section_path=current_path(),
                    clause_no=number,
                    text=line,
                    locator=locator,
                )
            )
            continue

        clauses.append(
            Clause(
                doc=doc_label,
                section_path=current_path(),
                clause_no="",
                text=line,
                locator=locator,
            )
        )
    return clauses


def _load_text(path: Path) -> List[Clause]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return parse_lines(text.splitlines(), path.name)


def _load_docx(path: Path) -> List[Clause]:
    import docx  # python-docx

    document = docx.Document(str(path))
    clauses: List[Clause] = []
    section_stack: List[Tuple[int, str]] = []

    def current_path() -> str:
        return " > ".join(title for _, title in section_stack)

    def push_heading(depth: int, title: str) -> None:
        while section_stack and section_stack[-1][0] >= depth:
            section_stack.pop()
        section_stack.append((depth, title))

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style.startswith("Heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            push_heading(int(digits) if digits else 1, text)
            continue
        numbered = _NUMBERED_RE.match(text)
        clause_no = ""
        if numbered and not _looks_like_heading(
            numbered.group(1), numbered.group(2).strip()
        ):
            clause_no = numbered.group(1)
        clauses.append(
            Clause(
                doc=path.name,
                section_path=current_path(),
                clause_no=clause_no,
                text=text,
                locator="第%d段" % index,
            )
        )

    for t_index, table in enumerate(document.tables, start=1):
        for r_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            text = " | ".join(cell for cell in cells if cell)
            if not text:
                continue
            clauses.append(
                Clause(
                    doc=path.name,
                    section_path=current_path(),
                    clause_no="",
                    text=text,
                    locator="表格%d第%d行" % (t_index, r_index),
                )
            )
    return clauses


def _load_pdf(path: Path) -> List[Clause]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    clauses: List[Clause] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        page_clauses = parse_lines(
            text.splitlines(), path.name, locator_fmt="第%d行"
        )
        for clause in page_clauses:
            clause.locator = "第%d页%s" % (page_no, clause.locator)
        clauses.extend(page_clauses)
    return clauses


def load_document(path) -> List[Clause]:
    """Load *path* and extract its clauses."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return _load_docx(file_path)
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix in {".txt", ".md", ".text"}:
        return _load_text(file_path)
    raise ValueError(
        "Unsupported file type %r; supported: %s"
        % (suffix, ", ".join(sorted(SUPPORTED_SUFFIXES)))
    )
