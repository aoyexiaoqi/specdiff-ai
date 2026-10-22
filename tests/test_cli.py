from openpyxl import load_workbook

from specdiff_ai.cli import main


def _write(path, text):
    path.write_text(text, encoding="utf-8")
    return str(path)


def test_cli_end_to_end(tmp_path):
    old = _write(tmp_path / "old.txt", "2 参数\n2.1 电流应为 8.5 kA。\n")
    new = _write(tmp_path / "new.txt", "2 参数\n2.1 电流应为 9.2 kA。\n")
    out = tmp_path / "result.xlsx"
    md = tmp_path / "result.md"
    assert main([old, new, "--out", str(out), "--md", str(md)]) == 0
    assert out.is_file() and md.is_file()
    sheet = load_workbook(str(out))["差异矩阵"]
    assert sheet.cell(row=2, column=2).value == "冲突"


def test_cli_missing_arguments():
    assert main([]) == 2


def test_cli_missing_file(tmp_path):
    assert main([str(tmp_path / "nope.txt"), str(tmp_path / "nope2.txt")]) == 2


def test_cli_demo(tmp_path):
    demo_dir = tmp_path / "demo"
    assert main(["--demo", "--demo-dir", str(demo_dir)]) == 0
    assert (demo_dir / "old_spec.txt").is_file()
    assert (demo_dir / "new_spec.txt").is_file()
    xlsx = demo_dir / "diff.xlsx"
    md = demo_dir / "diff.md"
    assert xlsx.is_file() and md.is_file()

    sheet = load_workbook(str(xlsx))["差异矩阵"]
    by_status = {}
    for row in range(2, sheet.max_row + 1):
        by_status.setdefault(sheet.cell(row=row, column=2).value, 0)
        by_status[sheet.cell(row=row, column=2).value] += 1
    # Planted differences: 2 conflicts, 1 modified, 1 added, 1 deleted, 6 unchanged
    assert by_status["冲突"] == 2
    assert by_status["修改"] == 1
    assert by_status["新增"] == 1
    assert by_status["删除"] == 1
    assert by_status["未变"] == 6
    assert "9.2 kA" in md.read_text(encoding="utf-8")
