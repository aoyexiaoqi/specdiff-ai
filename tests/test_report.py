from openpyxl import load_workbook

from specdiff_ai.diff import compare
from specdiff_ai.extract import parse_lines
from specdiff_ai.report import write_excel, write_markdown

OLD = ["2 参数", "2.1 电流应为 8.5 kA。"]
NEW = ["2 参数", "2.1 电流应为 9.2 kA。", "2.2 公差为 ±0.5 mm。"]


def _items():
    return compare(parse_lines(OLD, "old.txt"), parse_lines(NEW, "new.txt"))


def test_excel_matrix_contents(tmp_path):
    out = write_excel(_items(), tmp_path / "diff.xlsx", "old.txt", "new.txt")
    workbook = load_workbook(str(out))
    assert workbook.sheetnames == ["汇总", "差异矩阵"]
    sheet = workbook["差异矩阵"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:4] == ["序号", "状态", "章节", "条款号"]
    statuses = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
    assert statuses == ["冲突", "新增"]
    assert sheet.cell(row=2, column=7).value == "8.5 kA → 9.2 kA"
    assert "old.txt" in sheet.cell(row=2, column=8).value


def test_excel_summary_counts(tmp_path):
    out = write_excel(_items(), tmp_path / "diff.xlsx", "old.txt", "new.txt")
    workbook = load_workbook(str(out))
    summary = workbook["汇总"]
    values = {
        summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
        for r in range(1, summary.max_row + 1)
    }
    assert values["冲突"] == 1
    assert values["新增"] == 1
    assert values["旧版文档"] == "old.txt"


def test_markdown_report_contents(tmp_path):
    out = write_markdown(_items(), tmp_path / "diff.md", "old.txt", "new.txt")
    text = out.read_text(encoding="utf-8")
    assert "# 技术协议差异报告" in text
    assert "## 冲突（1）" in text
    assert "## 新增（1）" in text
    assert "8.5 kA → 9.2 kA" in text
    assert "old.txt" in text and "new.txt" in text
