"""Report writers: Excel difference matrix and Markdown report."""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from .diff import (
    ALL_STATUSES,
    STATUS_ADDED,
    STATUS_CONFLICT,
    STATUS_DELETED,
    STATUS_MODIFIED,
    DiffItem,
    count_by_status,
)

_STATUS_FILL = {
    STATUS_CONFLICT: "F8CBAD",
    STATUS_ADDED: "C6EFCE",
    STATUS_DELETED: "D9D9D9",
    STATUS_MODIFIED: "FFE699",
}
_STATUS_FONT = {
    STATUS_CONFLICT: "9C0006",
    STATUS_ADDED: "006100",
    STATUS_DELETED: "404040",
    STATUS_MODIFIED: "7F6000",
}


def write_excel(items: Sequence[DiffItem], path, old_name: str, new_name: str) -> Path:
    """Write the Excel difference matrix (汇总 sheet + 差异矩阵 sheet)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out = Path(path)
    counts = count_by_status(items)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary["A1"] = "SpecDiff AI 差异分析汇总"
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["旧版文档", old_name])
    summary.append(["新版文档", new_name])
    summary.append(["条款总数（新版顺序+删除项）", len(items)])
    summary.append([])
    summary.append(["状态", "数量"])
    for cell in summary[6]:
        cell.font = Font(bold=True)
    for status in ALL_STATUSES:
        summary.append([status, counts[status]])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 60

    sheet = workbook.create_sheet("差异矩阵")
    headers = [
        "序号", "状态", "章节", "条款号", "旧版内容", "新版内容",
        "参数变化", "旧版出处", "新版出处",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for idx, item in enumerate(items, start=1):
        old_evidence = "%s %s" % (item.old_doc, item.old_locator) if item.old_doc else ""
        new_evidence = "%s %s" % (item.new_doc, item.new_locator) if item.new_doc else ""
        sheet.append(
            [
                idx,
                item.status,
                item.section_path,
                item.clause_no,
                item.old_text,
                item.new_text,
                item.param_changes,
                old_evidence.strip(),
                new_evidence.strip(),
            ]
        )
        status_cell = sheet.cell(row=idx + 1, column=2)
        if item.status in _STATUS_FILL:
            status_cell.fill = PatternFill("solid", fgColor=_STATUS_FILL[item.status])
            status_cell.font = Font(color=_STATUS_FONT[item.status], bold=True)

    widths = [6, 8, 30, 10, 55, 55, 30, 22, 22]
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=2, max_col=9):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    workbook.save(str(out))
    return out


def _md_escape(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def write_markdown(items: Sequence[DiffItem], path, old_name: str, new_name: str) -> Path:
    """Write the Markdown difference report."""
    out = Path(path)
    counts = count_by_status(items)

    lines = []
    lines.append("# 技术协议差异报告 / Technical Agreement Diff Report")
    lines.append("")
    lines.append("- 旧版文档：`%s`" % old_name)
    lines.append("- 新版文档：`%s`" % new_name)
    lines.append("- 条款总数：%d" % len(items))
    lines.append("")
    lines.append("## 汇总 / Summary")
    lines.append("")
    lines.append("| 状态 | 数量 |")
    lines.append("| --- | --- |")
    for status in ALL_STATUSES:
        lines.append("| %s | %d |" % (status, counts[status]))
    lines.append("")

    for status in ALL_STATUSES:
        group = [item for item in items if item.status == status]
        if not group:
            continue
        lines.append("## %s（%d）" % (status, len(group)))
        lines.append("")
        for item in group:
            title = item.section_path or "（无章节）"
            if item.clause_no:
                title += " / 条款 " + item.clause_no
            lines.append("### [%s] %s" % (status, _md_escape(title)))
            lines.append("")
            if item.old_text:
                lines.append(
                    "- 旧版（%s %s）：%s"
                    % (item.old_doc, item.old_locator, _md_escape(item.old_text))
                )
            if item.new_text:
                lines.append(
                    "- 新版（%s %s）：%s"
                    % (item.new_doc, item.new_locator, _md_escape(item.new_text))
                )
            if item.param_changes:
                lines.append("- 参数变化：%s" % item.param_changes)
            lines.append("")

    lines.append("---")
    lines.append("")
    lines.append(
        "本报告由 SpecDiff AI 基于确定性规则生成；所有结论均附原文出处，"
        "请以源文档为准进行人工复核。"
    )
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out
